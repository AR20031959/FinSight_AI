import io
import datetime
from typing import Dict, Any, List

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import PieChart, BarChart, Reference
from openpyxl.utils import get_column_letter

from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, HRFlowable, Image as RLImage
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.pdfgen import canvas


class NumberedCanvas(canvas.Canvas):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_decorations(num_pages)
            super().showPage()
        super().save()

    def draw_page_decorations(self, page_count):
        self.saveState()
        self.setFont("Helvetica", 8.5)
        self.setFillColor(colors.HexColor("#64748B"))
        
        # Header border & text
        self.setStrokeColor(colors.HexColor("#CBD5E1"))
        self.setLineWidth(0.5)
        self.line(36, 756, 576, 756)
        self.drawString(36, 762, "FinSight AI - Executive Financial Decision Intelligence Platform")
        self.drawRightString(576, 762, "Confidential Financial Report")

        # Footer border & page number
        self.line(36, 45, 576, 45)
        self.drawString(36, 32, "Generated automatically by FinSight AI Engine")
        self.drawRightString(576, 32, f"Page {self._pageNumber} of {page_count}")
        self.restoreState()


def _generate_matplotlib_charts(summary: Dict[str, Any]) -> tuple[bytes, bytes]:
    """
    Generates two high-resolution Matplotlib PNG charts (Category Donut & Trend Timeline Bar)
    for embedding into PDF and PNG infographic reports.
    """
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    import numpy as np

    # 1. Category Pie / Donut Chart
    cat_breakdown = summary.get("category_breakdown", {})
    fig1, ax1 = plt.subplots(figsize=(5.2, 3.0), facecolor="#1E293B")
    ax1.set_facecolor("#1E293B")

    if cat_breakdown and any(float(v) > 0 for v in cat_breakdown.values()):
        cats = list(cat_breakdown.keys())[:6]
        vals = [float(cat_breakdown[k]) for k in cats]
        palette = ['#EF4444', '#F59E0B', '#8B5CF6', '#06B6D4', '#10B981', '#3B82F6']

        wedges, texts, autotexts = ax1.pie(
            vals, 
            labels=cats, 
            autopct='%1.1f%%',
            startangle=140,
            colors=palette[:len(cats)],
            textprops=dict(color="white", fontsize=8),
            wedgeprops=dict(width=0.45, edgecolor='#0F172A')
        )
        for autotext in autotexts:
            autotext.set_color('white')
            autotext.set_weight('bold')
            autotext.set_fontsize(7.5)
        ax1.set_title("Category Allocation Breakdown", color="white", fontsize=10, fontweight="bold", pad=10)
    else:
        ax1.text(0.5, 0.5, "No Category Expense Data Recorded", color="#94A3B8", fontsize=9.5, ha="center", va="center")
        ax1.set_title("Category Allocation Breakdown", color="white", fontsize=10, fontweight="bold", pad=10)
        ax1.axis('off')

    plt.tight_layout()
    buf1 = io.BytesIO()
    plt.savefig(buf1, format="png", dpi=140, facecolor=fig1.get_facecolor(), edgecolor='none')
    plt.close(fig1)
    buf1.seek(0)
    pie_bytes = buf1.getvalue()

    # 2. Period Cashflow Trend Bar Chart
    trend_data = summary.get("monthly_trend", [])
    if not trend_data and summary.get("yearly_trend"):
        trend_data = summary.get("yearly_trend", [])

    fig2, ax2 = plt.subplots(figsize=(5.2, 3.0), facecolor="#1E293B")
    ax2.set_facecolor("#1E293B")

    if trend_data:
        x_labels = [str(item.get("month") or item.get("year") or "") for item in trend_data]
        inc_vals = [float(item.get("income") or 0.0) for item in trend_data]
        exp_vals = [float(item.get("expense") or 0.0) for item in trend_data]

        x_indices = np.arange(len(x_labels))
        bar_w = 0.35

        ax2.bar(x_indices - bar_w/2, inc_vals, width=bar_w, label="Income", color="#10B981")
        ax2.bar(x_indices + bar_w/2, exp_vals, width=bar_w, label="Expense", color="#EF4444")
        ax2.set_xticks(x_indices)
        ax2.set_xticklabels(x_labels, rotation=20, color="white", fontsize=8)
        ax2.set_title("Period Cashflow Trend Timeline", color="white", fontsize=10, fontweight="bold", pad=10)
        ax2.tick_params(colors="white", labelsize=8)
        ax2.legend(facecolor="#1E293B", edgecolor="#475569", labelcolor="white", fontsize=7.5)
        for spine in ax2.spines.values():
            spine.set_color("#475569")
        ax2.grid(axis='y', linestyle='--', alpha=0.3, color='#64748B')
    else:
        ax2.text(0.5, 0.5, "No Period Trend Timeline Data", color="#94A3B8", fontsize=9.5, ha="center", va="center")
        ax2.set_title("Period Cashflow Trend Timeline", color="white", fontsize=10, fontweight="bold", pad=10)
        ax2.axis('off')

    plt.tight_layout()
    buf2 = io.BytesIO()
    plt.savefig(buf2, format="png", dpi=140, facecolor=fig2.get_facecolor(), edgecolor='none')
    plt.close(fig2)
    buf2.seek(0)
    trend_bytes = buf2.getvalue()

    return pie_bytes, trend_bytes


def _generate_spider_chart(summary: Dict[str, Any]) -> bytes:
    """
    Generates a high-resolution Matplotlib Polar Spider / Radar Chart
    representing multi-axis financial health pillars.
    """
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    import numpy as np

    categories = ['Savings', 'Investments', 'Expense Ctrl', 'Debt Safety', 'Runway', 'Health Score']
    N = len(categories)

    savings_rate = float(summary.get("savings_rate", 0.0) or 0.0)
    income = float(summary.get("monthly_income", 0.0) or 0.0)
    expense = float(summary.get("monthly_expense", 0.0) or 0.0)
    investments = float(summary.get("total_investments", 0.0) or 0.0)
    health = float(summary.get("health_score", 0.0) or 0.0)

    v1 = min(100.0, max(10.0, (savings_rate / 50.0) * 100.0))
    inv_ratio = (investments / income * 100.0) if income > 0 else (50.0 if investments > 0 else 10.0)
    v2 = min(100.0, max(10.0, (inv_ratio / 30.0) * 100.0))
    exp_ratio = (expense / income * 100.0) if income > 0 else 100.0
    v3 = min(100.0, max(15.0, 100.0 - (exp_ratio * 0.7)))
    v4 = 85.0
    liquid = float(summary.get("savings", 0.0) or 0.0) + investments
    runway = (liquid / max(expense, 1000.0))
    v5 = min(100.0, max(10.0, (runway / 6.0) * 100.0))
    v6 = min(100.0, max(15.0, health if health > 0 else 25.0))

    values = [v1, v2, v3, v4, v5, v6]
    benchmark = [75.0, 70.0, 80.0, 85.0, 80.0, 85.0]

    values += values[:1]
    benchmark += benchmark[:1]
    angles = [n / float(N) * 2 * np.pi for n in range(N)]
    angles += angles[:1]

    fig, ax = plt.subplots(figsize=(5.2, 3.0), subplot_kw=dict(polar=True), facecolor="#1E293B")
    ax.set_facecolor("#1E293B")
    ax.set_theta_offset(np.pi / 2)
    ax.set_theta_direction(-1)

    ax.set_xticks(angles[:-1])
    ax.set_xticklabels(categories, color="white", size=8, weight="bold")
    ax.set_rlabel_position(0)
    ax.tick_params(colors="white", labelsize=7)
    ax.grid(color="#334155", linestyle="--", alpha=0.6)

    ax.plot(angles, benchmark, color="#10B981", linewidth=1.5, linestyle="--", label="Target")
    ax.fill(angles, benchmark, color="#10B981", alpha=0.1)

    ax.plot(angles, values, color="#06B6D4", linewidth=2.2, label="Current")
    ax.fill(angles, values, color="#06B6D4", alpha=0.25)

    ax.set_ylim(0, 100)
    ax.set_title("Financial Health Radar", color="white", fontsize=10, fontweight="bold", pad=12)
    ax.legend(loc="lower right", bbox_to_anchor=(1.25, -0.1), facecolor="#1E293B", edgecolor="#475569", labelcolor="white", fontsize=7)

    plt.tight_layout()
    buf = io.BytesIO()
    plt.savefig(buf, format="png", dpi=140, facecolor=fig.get_facecolor(), edgecolor='none')
    plt.close(fig)
    buf.seek(0)
    return buf.getvalue()


def generate_pdf_report(
    user_name: str, 
    transactions: List[Dict[str, Any]], 
    summary: Dict[str, Any],
    start_date: str = None,
    end_date: str = None
) -> bytes:
    """
    Generates an executive, beautifully formatted PDF report containing metric cards,
    embedded high-res Matplotlib visual charts, health score insights, and clean paginated tables.
    """
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        leftMargin=36,
        rightMargin=36,
        topMargin=54,
        bottomMargin=54
    )
    story = []

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        'DocTitle', parent=styles['Heading1'],
        fontName='Helvetica-Bold', fontSize=20, leading=24,
        textColor=colors.HexColor('#0F172A'), spaceAfter=4
    )
    subtitle_style = ParagraphStyle(
        'DocSubtitle', parent=styles['Normal'],
        fontName='Helvetica', fontSize=9.5, leading=13,
        textColor=colors.HexColor('#64748B'), spaceAfter=10
    )
    h2_style = ParagraphStyle(
        'H2', parent=styles['Heading2'],
        fontName='Helvetica-Bold', fontSize=12, leading=16,
        textColor=colors.HexColor('#1E293B'), spaceBefore=10, spaceAfter=6
    )
    body_style = ParagraphStyle(
        'Body', parent=styles['Normal'],
        fontName='Helvetica', fontSize=8.5, leading=12,
        textColor=colors.HexColor('#334155')
    )
    table_cell_style = ParagraphStyle(
        'TableCell', parent=styles['Normal'],
        fontName='Helvetica', fontSize=8, leading=11,
        textColor=colors.HexColor('#1E293B')
    )

    if not start_date:
        start_date = "Beginning"
    if not end_date:
        end_date = datetime.datetime.now().strftime("%Y-%m-%d")

    date_range_str = f"Period: {start_date} to {end_date}"
    gen_time_str = datetime.datetime.now().strftime("%B %d, %Y at %H:%M")

    # Document Header
    story.append(Paragraph("FinSight AI - Executive Financial Intelligence Report", title_style))
    story.append(Paragraph(f"Account Holder: <b>{user_name}</b> | {date_range_str} | Generated: {gen_time_str}", subtitle_style))
    story.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor("#2563EB"), spaceAfter=12))

    # Calculate Totals
    tot_income = sum((t.get("amount") or 0.0) for t in transactions if str(t.get("type") or "").lower() == "income")
    tot_expense = sum((t.get("amount") or 0.0) for t in transactions if str(t.get("type") or "").lower() == "expense")
    tot_investment = sum((t.get("amount") or 0.0) for t in transactions if str(t.get("type") or "").lower() == "investment")
    net_surplus = tot_income - tot_expense - tot_investment

    if tot_income == 0 and summary.get("monthly_income"):
        tot_income = float(summary.get("monthly_income") or 0.0)
    if tot_expense == 0 and summary.get("monthly_expense"):
        tot_expense = float(summary.get("monthly_expense") or 0.0)
    if tot_investment == 0 and summary.get("total_investments"):
        tot_investment = float(summary.get("total_investments") or 0.0)
    if net_surplus == 0 and summary.get("savings"):
        net_surplus = float(summary.get("savings") or 0.0)

    # Executive Metric Cards Grid
    cards_data = [
        [
            Paragraph(f"<b>Period Income</b><br/><font size=12 color='#16A34A'><b>INR {tot_income:,.2f}</b></font>", body_style),
            Paragraph(f"<b>Period Expenses</b><br/><font size=12 color='#DC2626'><b>INR {tot_expense:,.2f}</b></font>", body_style),
            Paragraph(f"<b>Investments</b><br/><font size=12 color='#0284C7'><b>INR {tot_investment:,.2f}</b></font>", body_style),
            Paragraph(f"<b>Net Cashflow</b><br/><font size=12 color='#2563EB'><b>INR {net_surplus:,.2f}</b></font>", body_style),
        ]
    ]
    t_cards = Table(cards_data, colWidths=[135, 135, 135, 135])
    t_cards.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), colors.HexColor('#F8FAFC')),
        ('BOX', (0,0), (-1,-1), 1, colors.HexColor('#CBD5E1')),
        ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor('#E2E8F0')),
        ('TOPPADDING', (0,0), (-1,-1), 7),
        ('BOTTOMPADDING', (0,0), (-1,-1), 7),
        ('LEFTPADDING', (0,0), (-1,-1), 8),
        ('RIGHTPADDING', (0,0), (-1,-1), 8),
    ]))
    story.append(t_cards)
    story.append(Spacer(1, 10))

    # Financial Health Multi-Pillar Radar (Spider Chart) & AI Assessment
    story.append(Paragraph("Financial Health Multi-Pillar Radar & AI Assessment", h2_style))
    health_score = summary.get("health_score", 0.0)
    ai_insights = summary.get("ai_insights", [])
    insights_html = "<br/>".join([f"• {insight}" for insight in ai_insights[:3]]) if ai_insights else "• Multi-pillar financial metrics evaluated for this period."

    hs_text = (
        f"<b>Composite Health Score: <font color='#0284C7'>{health_score:.1f} / 100</font></b><br/><br/>"
        f"<b>Pillar Highlights:</b><br/>"
        f"{insights_html}<br/><br/>"
        f"<i>Multi-axis Spider Chart illustrates Savings, Investments, Expense Control, Debt Safety, Emergency Runway, and Health Score against benchmarks.</i>"
    )

    try:
        spider_bytes = _generate_spider_chart(summary)
        img_spider = RLImage(io.BytesIO(spider_bytes), width=265, height=150)

        t_hs_cell = Table([[Paragraph(hs_text, body_style)]], colWidths=[265])
        t_hs_cell.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,-1), colors.HexColor('#EFF6FF')),
            ('BOX', (0,0), (-1,-1), 1, colors.HexColor('#93C5FD')),
            ('TOPPADDING', (0,0), (-1,-1), 8),
            ('BOTTOMPADDING', (0,0), (-1,-1), 8),
            ('LEFTPADDING', (0,0), (-1,-1), 8),
            ('RIGHTPADDING', (0,0), (-1,-1), 8),
            ('VALIGN', (0,0), (-1,-1), 'TOP'),
        ]))

        radar_table = Table([[t_hs_cell, img_spider]], colWidths=[270, 270])
        radar_table.setStyle(TableStyle([
            ('ALIGN', (0,0), (-1,-1), 'CENTER'),
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
            ('BOTTOMPADDING', (0,0), (-1,-1), 0),
            ('TOPPADDING', (0,0), (-1,-1), 0),
            ('LEFTPADDING', (0,0), (-1,-1), 0),
            ('RIGHTPADDING', (0,0), (-1,-1), 0),
        ]))
        story.append(radar_table)
    except Exception:
        t_hs = Table([[Paragraph(hs_text, body_style)]], colWidths=[540])
        t_hs.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,-1), colors.HexColor('#EFF6FF')),
            ('BOX', (0,0), (-1,-1), 1, colors.HexColor('#93C5FD')),
            ('TOPPADDING', (0,0), (-1,-1), 7),
            ('BOTTOMPADDING', (0,0), (-1,-1), 7),
            ('LEFTPADDING', (0,0), (-1,-1), 10),
            ('RIGHTPADDING', (0,0), (-1,-1), 10),
        ]))
        story.append(t_hs)

    story.append(Spacer(1, 10))

    # Embedded High-Res Visual Analytics Charts (Category Donut & Cashflow Trend)
    try:
        pie_bytes, trend_bytes = _generate_matplotlib_charts(summary)
        img_pie = RLImage(io.BytesIO(pie_bytes), width=265, height=150)
        img_trend = RLImage(io.BytesIO(trend_bytes), width=265, height=150)

        charts_table = Table([[img_pie, img_trend]], colWidths=[270, 270])
        charts_table.setStyle(TableStyle([
            ('ALIGN', (0,0), (-1,-1), 'CENTER'),
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
            ('BOTTOMPADDING', (0,0), (-1,-1), 0),
            ('TOPPADDING', (0,0), (-1,-1), 0),
            ('LEFTPADDING', (0,0), (-1,-1), 0),
            ('RIGHTPADDING', (0,0), (-1,-1), 0),
        ]))
        story.append(Paragraph("Visual Analytics & Allocation Overview", h2_style))
        story.append(charts_table)
        story.append(Spacer(1, 12))
    except Exception:
        pass

    # Detailed Transactions Activity Ledger
    story.append(Paragraph(f"Activity Ledger ({len(transactions)} Records in Selected Period)", h2_style))

    headers = [
        Paragraph("<b>Date</b>", table_cell_style),
        Paragraph("<b>Title / Description</b>", table_cell_style),
        Paragraph("<b>Category</b>", table_cell_style),
        Paragraph("<b>Type</b>", table_cell_style),
        Paragraph("<b>Amount (INR)</b>", table_cell_style),
    ]
    tx_rows = [headers]

    for t in transactions:
        tx_type = str(t.get("type") or "").upper()
        amt = float(t.get("amount") or 0.0)
        amt_str = f"INR {amt:,.2f}"

        tx_rows.append([
            Paragraph(str(t.get("date") or ""), table_cell_style),
            Paragraph(str(t.get("title") or ""), table_cell_style),
            Paragraph(str(t.get("category") or ""), table_cell_style),
            Paragraph(tx_type, table_cell_style),
            Paragraph(amt_str, table_cell_style),
        ])

    if len(transactions) == 0:
        tx_rows.append([
            Paragraph(f"No transactions recorded between {start_date} and {end_date}", table_cell_style),
            Paragraph("", table_cell_style),
            Paragraph("", table_cell_style),
            Paragraph("", table_cell_style),
            Paragraph("", table_cell_style)
        ])

    t_txs = Table(tx_rows, colWidths=[70, 190, 105, 75, 100], repeatRows=1)
    t_txs.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#1E293B')),
        ('TEXTCOLOR', (0,0), (-1,0), colors.white),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#CBD5E1')),
        ('TOPPADDING', (0,0), (-1,-1), 5),
        ('BOTTOMPADDING', (0,0), (-1,-1), 5),
        ('LEFTPADDING', (0,0), (-1,-1), 5),
        ('RIGHTPADDING', (0,0), (-1,-1), 5),
        ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor('#F8FAFC')]),
    ]))
    story.append(t_txs)

    doc.build(story, canvasmaker=NumberedCanvas)
    buffer.seek(0)
    return buffer.getvalue()


def generate_excel_report(
    user_name: str, 
    transactions: List[Dict[str, Any]], 
    summary: Dict[str, Any],
    start_date: str = None,
    end_date: str = None
) -> bytes:
    """
    Generates a structured, multi-tab Excel workbook containing:
    1. Executive Dashboard tab with styled KPI cards & native OpenPyXL Bar/Pie charts.
    2. Detailed Activity Ledger tab with formatted currency, auto-fitted columns, and zebra striping.
    """
    wb = Workbook()

    # Define Colors & Fills
    title_font = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
    section_font = Font(name="Calibri", size=12, bold=True, color="1E293B")
    bold_font = Font(name="Calibri", size=11, bold=True)
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

    title_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    kpi_inc_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    kpi_exp_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
    kpi_inv_fill = PatternFill(start_color="CFFAFE", end_color="CFFAFE", fill_type="solid")
    kpi_sav_fill = PatternFill(start_color="DBEAFE", end_color="DBEAFE", fill_type="solid")
    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    if not start_date:
        start_date = "Beginning"
    if not end_date:
        end_date = datetime.datetime.now().strftime("%Y-%m-%d")

    # TAB 1: EXECUTIVE DASHBOARD
    ws_dash = wb.active
    ws_dash.title = "Executive Dashboard"
    ws_dash.views.sheetView[0].showGridLines = True

    # Title Banner
    ws_dash.merge_cells("A1:D1")
    title_cell = ws_dash["A1"]
    title_cell.value = "FinSight AI - Executive Financial Dashboard Report"
    title_cell.font = title_font
    title_cell.fill = title_fill
    title_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws_dash.row_dimensions[1].height = 32

    ws_dash["A2"] = f"Account Holder: {user_name}  |  Period: {start_date} to {end_date}  |  Generated: {datetime.datetime.now().strftime('%Y-%m-%d %H:%M')}"
    ws_dash["A2"].font = Font(name="Calibri", size=10, italic=True, color="64748B")

    # KPI Summary Cards
    income = float(summary.get("monthly_income") or sum((t.get("amount") or 0.0) for t in transactions if str(t.get("type") or "").lower() == "income"))
    expense = float(summary.get("monthly_expense") or sum((t.get("amount") or 0.0) for t in transactions if str(t.get("type") or "").lower() == "expense"))
    investments = float(summary.get("total_investments") or sum((t.get("amount") or 0.0) for t in transactions if str(t.get("type") or "").lower() == "investment"))
    savings = float(summary.get("savings") or (income - expense - investments))
    health_score = float(summary.get("health_score") or 0.0)

    ws_dash["A4"] = "KPI Financial Summary"
    ws_dash["A4"].font = section_font

    kpi_rows = [
        ("Period Income", income, kpi_inc_fill, "16A34A"),
        ("Period Expenses", expense, kpi_exp_fill, "DC2626"),
        ("Total Investments", investments, kpi_inv_fill, "0284C7"),
        ("Net Cashflow Savings", savings, kpi_sav_fill, "2563EB"),
        ("Financial Health Score", f"{health_score:.1f} / 100", None, "8B5CF6"),
    ]

    for idx, (label, val, fill_style, text_color) in enumerate(kpi_rows, start=5):
        ws_dash[f"A{idx}"] = label
        ws_dash[f"A{idx}"].font = bold_font
        ws_dash[f"B{idx}"] = val
        ws_dash[f"B{idx}"].font = Font(name="Calibri", size=11, bold=True, color=text_color)
        if isinstance(val, (int, float)):
            ws_dash[f"B{idx}"].number_format = '"INR "#,##0.00'
        if fill_style:
            ws_dash[f"A{idx}"].fill = fill_style
            ws_dash[f"B{idx}"].fill = fill_style
        ws_dash[f"A{idx}"].border = thin_border
        ws_dash[f"B{idx}"].border = thin_border

    # Category Allocation Table (for Chart Data Source)
    cat_breakdown = summary.get("category_breakdown", {})
    cats = list(cat_breakdown.keys())
    
    ws_dash["A12"] = "Category Allocation Data"
    ws_dash["A12"].font = section_font

    ws_dash["A13"] = "Category"
    ws_dash["B13"] = "Amount (INR)"
    ws_dash["A13"].font = header_font
    ws_dash["A13"].fill = header_fill
    ws_dash["B13"].font = header_font
    ws_dash["B13"].fill = header_fill

    current_row = 14
    if cats:
        for cat, val in cat_breakdown.items():
            ws_dash[f"A{current_row}"] = cat
            ws_dash[f"B{current_row}"] = float(val)
            ws_dash[f"B{current_row}"].number_format = '"INR "#,##0.00'
            ws_dash[f"A{current_row}"].border = thin_border
            ws_dash[f"B{current_row}"].border = thin_border
            current_row += 1
    else:
        ws_dash[f"A{current_row}"] = "No Categories"
        ws_dash[f"B{current_row}"] = 0.0
        current_row += 1

    cat_end_row = max(14, current_row - 1)

    # OpenPyXL Native PieChart for Category Breakdown
    if cats:
        try:
            pie = PieChart()
            pie.title = "Category Spending Allocation"
            labels_ref = Reference(ws_dash, min_col=1, min_row=14, max_row=cat_end_row)
            data_ref = Reference(ws_dash, min_col=2, min_row=13, max_row=cat_end_row)
            pie.add_data(data_ref, titles_from_data=True)
            pie.set_categories(labels_ref)
            pie.width = 15
            pie.height = 9
            ws_dash.add_chart(pie, "E4")
        except Exception:
            pass

    # Period Trend Data Table (for BarChart)
    trend_data = summary.get("monthly_trend", [])
    if not trend_data and summary.get("yearly_trend"):
        trend_data = summary.get("yearly_trend", [])

    trend_start_row = current_row + 2
    ws_dash[f"A{trend_start_row}"] = "Period Trend Data"
    ws_dash[f"A{trend_start_row}"].font = section_font

    t_hdr_row = trend_start_row + 1
    ws_dash[f"A{t_hdr_row}"] = "Period"
    ws_dash[f"B{t_hdr_row}"] = "Income"
    ws_dash[f"C{t_hdr_row}"] = "Expense"
    for col_c in ["A", "B", "C"]:
        ws_dash[f"{col_c}{t_hdr_row}"].font = header_font
        ws_dash[f"{col_c}{t_hdr_row}"].fill = header_fill

    t_data_row = t_hdr_row + 1
    if trend_data:
        for item in trend_data:
            ws_dash[f"A{t_data_row}"] = str(item.get("month") or item.get("year") or "")
            ws_dash[f"B{t_data_row}"] = float(item.get("income") or 0.0)
            ws_dash[f"C{t_data_row}"] = float(item.get("expense") or 0.0)
            ws_dash[f"B{t_data_row}"].number_format = '"INR "#,##0.00'
            ws_dash[f"C{t_data_row}"].number_format = '"INR "#,##0.00'
            for col_c in ["A", "B", "C"]:
                ws_dash[f"{col_c}{t_data_row}"].border = thin_border
            t_data_row += 1
    else:
        ws_dash[f"A{t_data_row}"] = "Current"
        ws_dash[f"B{t_data_row}"] = income
        ws_dash[f"C{t_data_row}"] = expense
        t_data_row += 1

    t_end_row = t_data_row - 1

    # OpenPyXL Native BarChart for Period Trend
    try:
        bar = BarChart()
        bar.type = "col"
        bar.style = 10
        bar.title = "Cashflow Income vs Expense Trend"
        bar.y_axis.title = "Amount (INR)"
        bar.x_axis.title = "Period"
        data_ref = Reference(ws_dash, min_col=2, min_row=t_hdr_row, max_col=3, max_row=t_end_row)
        cats_ref = Reference(ws_dash, min_col=1, min_row=t_hdr_row + 1, max_row=t_end_row)
        bar.add_data(data_ref, titles_from_data=True)
        bar.set_categories(cats_ref)
        bar.width = 15
        bar.height = 9
        ws_dash.add_chart(bar, "E20")
    except Exception:
        pass

    # Embed Multi-Pillar Spider / Radar Chart into Excel Dashboard
    try:
        from openpyxl.drawing.image import Image as OpenPyXLImage
        spider_bytes = _generate_spider_chart(summary)
        spider_img = OpenPyXLImage(io.BytesIO(spider_bytes))
        spider_img.width = 450
        spider_img.height = 250
        ws_dash.add_image(spider_img, "E35")
    except Exception:
        pass

    # TAB 2: DETAILED ACTIVITY LEDGER
    ws_ledger = wb.create_sheet(title="Activity Ledger")
    ws_ledger.views.sheetView[0].showGridLines = True

    ledger_headers = ["ID", "Date", "Title / Description", "Category", "Type", "Amount (INR)", "Merchant"]
    ws_ledger.append(ledger_headers)

    ws_ledger.row_dimensions[1].height = 26
    for cell in ws_ledger[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for idx, t in enumerate(transactions, start=2):
        row_values = [
            t.get("id", ""),
            t.get("date", ""),
            t.get("title", ""),
            t.get("category", ""),
            str(t.get("type") or "").upper(),
            float(t.get("amount") or 0.0),
            t.get("merchant", "")
        ]
        ws_ledger.append(row_values)
        ws_ledger.row_dimensions[idx].height = 20
        
        # Formatting & Borders
        for col_idx, val in enumerate(row_values, start=1):
            cell = ws_ledger.cell(row=idx, column=col_idx)
            cell.border = thin_border
            if idx % 2 == 1:
                cell.fill = zebra_fill
            if col_idx == 6:  # Amount Column
                cell.number_format = '"INR "#,##0.00'
                cell.alignment = Alignment(horizontal="right", vertical="center")

    # Auto-fit Column Widths for both sheets
    for ws in [ws_dash, ws_ledger]:
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def generate_image_report(
    user_name: str, 
    transactions: List[Dict[str, Any]], 
    summary: Dict[str, Any], 
    start_date: str = None, 
    end_date: str = None
) -> bytes:
    """
    Generates a high-resolution visual financial infographic image report (PNG format)
    with dynamic height, metric card containers, embedded Matplotlib visual charts,
    and a padded, non-overlapping transaction ledger.
    """
    from PIL import Image, ImageDraw, ImageFont

    txs_to_show = transactions[:15]

    # Calculate exact dynamic canvas height to eliminate overlap
    header_h = 130
    cards_h = 230
    health_h = 320
    chart_h = 380
    table_header_h = 50
    table_rows_h = max(1, len(txs_to_show)) * 48
    footer_h = 60
    
    total_canvas_h = 40 + header_h + 20 + cards_h + 20 + health_h + 20 + chart_h + 20 + table_header_h + table_rows_h + footer_h + 40

    img = Image.new("RGB", (1200, total_canvas_h), color="#0F172A")
    draw = ImageDraw.Draw(img)

    # Try loading custom fonts, fallback to default
    try:
        font_title = ImageFont.truetype("arial.ttf", 22)
        font_sub = ImageFont.truetype("arial.ttf", 13)
        font_card_lbl = ImageFont.truetype("arial.ttf", 12)
        font_card_val = ImageFont.truetype("arial.ttf", 20)
        font_table_hdr = ImageFont.truetype("arial.ttf", 12)
        font_table_cell = ImageFont.truetype("arial.ttf", 11)
    except Exception:
        font_title = font_sub = font_card_lbl = font_card_val = font_table_hdr = font_table_cell = ImageFont.load_default()

    y_cursor = 40

    # 1. Header Card
    draw.rectangle([40, y_cursor, 1160, y_cursor + header_h], fill="#1E293B", outline="#3B82F6", width=2)
    draw.text((60, y_cursor + 20), "FINSIGHT AI - EXECUTIVE FINANCIAL INFOGRAPHIC REPORT", fill="#38BDF8", font=font_title)

    period_str = f"{start_date or 'Beginning'} to {end_date or 'Present'}"
    draw.text((60, y_cursor + 60), f"Account Holder: {user_name}  |  Report Period: {period_str}", fill="#94A3B8", font=font_sub)
    draw.text((60, y_cursor + 90), f"Generated: {datetime.datetime.now().strftime('%d %B %Y, %I:%M %p')}", fill="#64748B", font=font_sub)

    y_cursor += header_h + 20

    # 2. Metric Cards Grid
    income = float(summary.get("monthly_income") or sum((t.get("amount") or 0.0) for t in transactions if str(t.get("type") or "").lower() == "income"))
    expense = float(summary.get("monthly_expense") or sum((t.get("amount") or 0.0) for t in transactions if str(t.get("type") or "").lower() == "expense"))
    invest = float(summary.get("total_investments") or sum((t.get("amount") or 0.0) for t in transactions if str(t.get("type") or "").lower() == "investment"))
    savings = float(summary.get("savings") or (income - expense - invest))
    health = float(summary.get("health_score") or 0.0)

    cards = [
        ("PERIOD INCOME", f"INR {income:,.2f}", "#10B981", (40, y_cursor, 580, y_cursor + 100)),
        ("PERIOD EXPENSES", f"INR {expense:,.2f}", "#EF4444", (620, y_cursor, 1160, y_cursor + 100)),
        ("INVESTMENTS", f"INR {invest:,.2f}", "#06B6D4", (40, y_cursor + 110, 580, y_cursor + 210)),
        ("NET SURPLUS SAVINGS", f"INR {savings:,.2f}", "#2563EB", (620, y_cursor + 110, 1160, y_cursor + 210)),
    ]

    for label, val, color, bounds in cards:
        draw.rectangle(bounds, fill="#1E293B", outline=color, width=2)
        draw.text((bounds[0] + 20, bounds[1] + 18), label, fill="#94A3B8", font=font_card_lbl)
        draw.text((bounds[0] + 20, bounds[1] + 50), val, fill=color, font=font_card_val)

    y_cursor += cards_h + 20

    # 3. Health Score & AI Insight Card (Left) + Multi-Pillar Spider Radar Chart (Right)
    draw.rectangle([40, y_cursor, 580, y_cursor + health_h], fill="#1E293B", outline="#8B5CF6", width=2)
    draw.text((60, y_cursor + 20), f"FINANCIAL HEALTH SCORE: {health:.1f} / 100", fill="#A855F7", font=font_card_val)

    ai_insights = summary.get("ai_insights", [])
    y_in = y_cursor + 65
    insights_to_show = ai_insights[:4] if ai_insights else ["Multi-pillar financial evaluation complete for this period."]
    for insight in insights_to_show:
        ins_str = f"• {insight}"
        if len(ins_str) > 52:
            ins_str = ins_str[:49] + "..."
        draw.text((60, y_in), ins_str, fill="#E2E8F0", font=font_sub)
        y_in += 30

    draw.text((60, y_cursor + health_h - 45), "Spider Radar reflects 6 core health pillars vs benchmark.", fill="#94A3B8", font=font_card_lbl)

    try:
        spider_bytes = _generate_spider_chart(summary)
        chart_spider_img = Image.open(io.BytesIO(spider_bytes)).resize((540, health_h))
        img.paste(chart_spider_img, (620, y_cursor))
    except Exception:
        draw.rectangle([620, y_cursor, 1160, y_cursor + health_h], fill="#1E293B", outline="#475569", width=1)
        draw.text((650, y_cursor + 140), "Financial Health Radar Chart", fill="#94A3B8", font=font_title)

    y_cursor += health_h + 20

    # 4. Embedded Dual Matplotlib Charts (Side-by-side)
    chart_y_start = y_cursor
    try:
        pie_bytes, trend_bytes = _generate_matplotlib_charts(summary)

        chart_pie_img = Image.open(io.BytesIO(pie_bytes)).resize((540, 360))
        chart_trend_img = Image.open(io.BytesIO(trend_bytes)).resize((540, 360))

        img.paste(chart_pie_img, (40, chart_y_start))
        img.paste(chart_trend_img, (620, chart_y_start))
    except Exception:
        draw.rectangle([40, chart_y_start, 1160, chart_y_start + chart_h], fill="#1E293B", outline="#475569", width=1)
        draw.text((60, chart_y_start + 160), "Visual Analytics Charts Overview", fill="#94A3B8", font=font_title)

    y_cursor += chart_h + 20

    # 5. Transaction Table Header
    draw.rectangle([40, y_cursor, 1160, y_cursor + table_header_h], fill="#2563EB")
    draw.text((60, y_cursor + 16), "DATE", fill="white", font=font_table_hdr)
    draw.text((200, y_cursor + 16), "DESCRIPTION / TITLE", fill="white", font=font_table_hdr)
    draw.text((650, y_cursor + 16), "CATEGORY", fill="white", font=font_table_hdr)
    draw.text((850, y_cursor + 16), "TYPE", fill="white", font=font_table_hdr)
    draw.text((980, y_cursor + 16), "AMOUNT (INR)", fill="white", font=font_table_hdr)

    y_cursor += table_header_h

    # 6. Table Rows with Zebra Striping
    if not txs_to_show:
        draw.rectangle([40, y_cursor, 1160, y_cursor + 48], fill="#1E293B")
        draw.text((60, y_cursor + 16), "No transactions recorded for this period.", fill="#94A3B8", font=font_table_cell)
        y_cursor += 48
    else:
        for idx, t in enumerate(txs_to_show):
            row_bg = "#1E293B" if idx % 2 == 0 else "#0F172A"
            draw.rectangle([40, y_cursor, 1160, y_cursor + 45], fill=row_bg)

            t_type = str(t.get("type") or "").lower()
            amt = float(t.get("amount") or 0.0)
            amt_color = "#10B981" if t_type == "income" else ("#06B6D4" if t_type in ("investment", "savings") else "#EF4444")

            title_str = str(t.get("title") or "")
            if len(title_str) > 34:
                title_str = title_str[:31] + "..."

            cat_str = str(t.get("category") or "")
            if len(cat_str) > 20:
                cat_str = cat_str[:17] + "..."

            draw.text((60, y_cursor + 14), str(t.get("date") or "")[:10], fill="#94A3B8", font=font_table_cell)
            draw.text((200, y_cursor + 14), title_str, fill="white", font=font_table_cell)
            draw.text((650, y_cursor + 14), cat_str, fill="#94A3B8", font=font_table_cell)
            draw.text((850, y_cursor + 14), t_type.upper(), fill="#94A3B8", font=font_table_cell)
            draw.text((980, y_cursor + 14), f"INR {amt:,.2f}", fill=amt_color, font=font_table_cell)

            y_cursor += 48

    y_cursor += 20
    # 7. Report Footer
    draw.line([40, y_cursor, 1160, y_cursor], fill="#334155", width=1)
    draw.text((60, y_cursor + 15), "FinSight AI Decision Intelligence Platform  |  Automated Confidential Report", fill="#64748B", font=font_sub)

    buf = io.BytesIO()
    img.save(buf, format="PNG")
    buf.seek(0)
    return buf.getvalue()
